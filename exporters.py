from __future__ import annotations

import csv
import shutil
from pathlib import Path

from openpyxl import Workbook

from models import AnalysisResult, UniqueUserRecord
from raw_comment_spool import RAW_COMMENT_COLUMNS, iter_raw_comment_rows
from utils import format_datetime, raw_comments_csv_path_for, report_path_for, report_to_text

EXCEL_MAX_ROWS = 1_048_576


class ExportError(RuntimeError):
    """Ошибка экспорта, которую можно показать пользователю без traceback."""


UNIQUE_USER_COLUMNS = [
    "user_id",
    "username",
    "profile_url",
    "first_name",
    "last_name",
    "is_bot",
    "first_seen_comment_date",
    "last_seen_comment_date",
    "occurrences_count",
    "source_channels_count",
]


def export_results(
    *,
    result: AnalysisResult,
    output_path: Path,
    export_format: str,
) -> list[Path]:
    output_path.parent.mkdir(parents=True, exist_ok=True)

    report_text = report_to_text(result.report)
    report_file = report_path_for(output_path)

    exported_files: list[Path] = []

    try:
        if export_format == "xlsx":
            _export_xlsx(result, output_path)
            exported_files.append(output_path)
        else:
            raw_comments_path = raw_comments_csv_path_for(output_path)
            _export_csv(result, output_path, raw_comments_path)
            exported_files.extend([output_path, raw_comments_path])

        report_file.write_text(report_text, encoding="utf-8")
        exported_files.append(report_file)
        return exported_files
    finally:
        if result.raw_comments_path is not None:
            result.raw_comments_path.unlink(missing_ok=True)


def _export_csv(
    result: AnalysisResult,
    unique_users_path: Path,
    raw_comments_path: Path,
) -> None:
    with unique_users_path.open("w", encoding="utf-8-sig", newline="") as file:
        writer = csv.DictWriter(file, fieldnames=UNIQUE_USER_COLUMNS)
        writer.writeheader()
        for record in result.unique_users:
            writer.writerow(_unique_user_row(record))

    if result.raw_comments_path is None:
        _write_empty_raw_comments_csv(raw_comments_path)
        return

    shutil.copyfile(result.raw_comments_path, raw_comments_path)


def _export_xlsx(result: AnalysisResult, output_path: Path) -> None:
    _validate_excel_row_count("unique_users", len(result.unique_users))
    _validate_excel_row_count("raw_comments", result.raw_comments_count)

    workbook = Workbook(write_only=True)

    unique_sheet = workbook.create_sheet("unique_users")
    unique_sheet.append(UNIQUE_USER_COLUMNS)
    for record in result.unique_users:
        row = _unique_user_row(record)
        unique_sheet.append([row[column] for column in UNIQUE_USER_COLUMNS])

    raw_sheet = workbook.create_sheet("raw_comments")
    raw_sheet.append(RAW_COMMENT_COLUMNS)
    for row in iter_raw_comment_rows(result.raw_comments_path):
        raw_sheet.append(_raw_comment_xlsx_values(row))

    workbook.save(output_path)


def _write_empty_raw_comments_csv(path: Path) -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as file:
        writer = csv.DictWriter(file, fieldnames=RAW_COMMENT_COLUMNS)
        writer.writeheader()


def _unique_user_row(record: UniqueUserRecord) -> dict[str, object]:
    return {
        "user_id": record.user_id,
        "username": record.username or "",
        "profile_url": record.profile_url,
        "first_name": record.first_name or "",
        "last_name": record.last_name or "",
        "is_bot": record.is_bot,
        "first_seen_comment_date": format_datetime(record.first_seen_comment_date),
        "last_seen_comment_date": format_datetime(record.last_seen_comment_date),
        "occurrences_count": record.occurrences_count,
        "source_channels_count": record.source_channels_count,
    }


def _raw_comment_xlsx_values(row: dict[str, str]) -> list[object]:
    return [
        _parse_int(row.get("user_id", "")),
        row.get("username", ""),
        row.get("profile_url", ""),
        row.get("first_name", ""),
        row.get("last_name", ""),
        _parse_bool(row.get("is_bot", "")),
        row.get("source_channel", ""),
        _parse_int(row.get("source_post_id", "")),
        row.get("comment_date", ""),
    ]


def _validate_excel_row_count(sheet_name: str, records_count: int) -> None:
    if records_count > EXCEL_MAX_ROWS - 1:
        raise ExportError(
            f"Лист '{sheet_name}' содержит {records_count} строк данных, "
            f"но XLSX поддерживает максимум {EXCEL_MAX_ROWS - 1}. "
            "Используйте --format csv или уменьшите лимиты анализа."
        )


def _parse_int(value: object) -> int | str:
    try:
        return int(str(value))
    except (TypeError, ValueError):
        return ""


def _parse_bool(value: object) -> bool | str:
    normalized = str(value).strip().lower()
    if normalized == "true":
        return True
    if normalized == "false":
        return False
    return ""
